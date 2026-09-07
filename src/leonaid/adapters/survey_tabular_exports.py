"""Deterministic, formula-safe CSV/XLSX from one immutable survey selection."""

import csv
import io
import json
import math
import re
import unicodedata
import zipfile
from datetime import datetime
from typing import Any
from xml.etree import ElementTree

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from leonaid.application.surveys.analysis_snapshot import AnalysisSnapshot
from leonaid.application.surveys.analysis import AggregateCount
from leonaid.application.surveys.export_rendering import (
    ExportProduct,
    SurveyExportArtifact,
    SurveyExportRenderError,
    SurveyExportSource,
    export_filename,
)

RENDER_VERSION = "survey-tabular-v1"
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
BASE_COLUMNS = [
    "participation_id",
    "revision",
    "status",
    "created_at",
    "completed_at",
    "last_page",
]


def json_value(value: Any) -> str:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        allow_nan=False,
    )


def formula_like(value: str) -> bool:
    rest = value
    while rest and (rest[0].isspace() or unicodedata.category(rest[0]).startswith("C")):
        rest = rest[1:]
    return bool(rest and rest[0] in "=+-@") or value.startswith(("\t", "\r", "\n"))


def csv_value(value: Any) -> Any:
    if isinstance(value, str) and (formula_like(value) or value.startswith("'")):
        return "'" + value
    if isinstance(value, bool):
        return "true" if value else "false"
    return value


def cell_value(value: Any) -> Any:
    if isinstance(value, (list, dict)):
        return json_value(value)
    if isinstance(value, float) and not math.isfinite(value):
        raise SurveyExportRenderError("non_finite_value")
    # Excel only retains 15 significant decimal digits for numeric cells.
    if (
        isinstance(value, int)
        and not isinstance(value, bool)
        and len(str(abs(value))) > 15
    ):
        return str(value)
    return value


def raw_value(value: Any) -> Any:
    if (
        isinstance(value, float)
        and math.isfinite(value)
        and float(format(value, ".15g")) != value
    ):
        return repr(value)
    return cell_value(value)


def value_type(answers: dict[str, Any], key: str) -> str:
    if key not in answers:
        return "missing"
    value = answers[key]
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, str):
        return "string"
    return "array" if isinstance(value, list) else "object"


def excel_text(value: str) -> str:
    # openpyxl otherwise silently truncates long strings to Excel's cell limit.
    if len(value) > 32767 or INVALID_XML.search(value):
        raise SurveyExportRenderError("cell_text_unrepresentable")
    return value


def append(sheet: Any, values: list[Any]) -> None:
    row = sheet.max_row + 1 if sheet.cell(1, 1).value is not None else 1
    for column, value in enumerate(values, 1):
        value = cell_value(value)
        if isinstance(value, str):
            value = excel_text(value)
        cell = sheet.cell(row, column, value)
        if isinstance(value, str):
            # Explicit strings prevent formula/hyperlink interpretation while preserving text.
            cell.data_type = "s"
            cell.quotePrefix = formula_like(value)
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def metadata(source: SurveyExportSource) -> dict[str, Any]:
    s = source.snapshot
    return {
        "format_version": RENDER_VERSION,
        "snapshot_id": s.id,
        "survey_id": s.surveyId,
        "survey_title": source.title,
        "version_id": s.filter.versionId,
        "version_number": s.versionNumber,
        "renderer_version": s.rendererVersion,
        "capability_profile": s.capabilityProfile,
        "captured_at": s.createdAt,
        "statuses_json": json_value(s.filter.statuses),
        "test_data": s.filter.isTest,
        "created_from_inclusive": s.filter.createdFrom,
        "created_before_exclusive": s.filter.createdBefore,
        "selected_participations": s.participationCount,
        "scope_in_progress": s.statusCounts.in_progress,
        "scope_partial": s.statusCounts.partial,
        "scope_completed": s.statusCounts.completed,
        "denominators": "Valid answered participants per question; valid cells per matrix row. Missing is not zero. Multiselect totals can exceed 100%.",
        "raw_encoding": "Structured values use canonical JSON. Each q:ID:type preserves missing/null/string/number/boolean/array/object, including empty strings and exact high-precision numeric text. Matrix row columns supplement the full JSON value.",
        "text_safety": "CSV prefixes formula-like and apostrophe-leading text with one apostrophe; remove that one prefix when decoding strings. XLSX uses explicit strings; unrepresentable cells fail without truncation.",
    }


def response_rows(source: SurveyExportSource) -> tuple[list[str], list[list[Any]]]:
    s = source.snapshot
    if source.responses is None or len(source.responses) != s.participationCount:
        raise SurveyExportRenderError("raw_selection_required")
    if len({r.participationId for r in source.responses}) != len(source.responses):
        raise SurveyExportRenderError("duplicate_participation")
    if any(r.status not in s.filter.statuses for r in source.responses):
        raise SurveyExportRenderError("response_status_outside_selection")
    columns = list(BASE_COLUMNS)
    for q in s.questions:
        columns.extend([f"q:{q.questionId}", f"q:{q.questionId}:type"])
        columns.extend(f"q:{q.questionId}:row:{r.rowId}" for r in q.matrixRows)
    rows = []
    for response in sorted(
        source.responses, key=lambda r: (r.createdAt, r.participationId)
    ):
        values = [
            response.participationId,
            response.revision,
            response.status,
            response.createdAt,
            response.completedAt,
            response.currentPage,
        ]
        for q in s.questions:
            answer = response.answers.get(q.questionId)
            values.extend(
                [raw_value(answer), value_type(response.answers, q.questionId)]
            )
            values.extend(
                raw_value(answer.get(r.rowId)) if isinstance(answer, dict) else None
                for r in q.matrixRows
            )
        rows.append(values)
    return columns, rows


def catalogue(workbook: Any, source: SurveyExportSource) -> None:
    sheet = workbook.create_sheet("Questions")
    append(
        sheet, ["record_type", "question_id", "title", "kind", "value_json", "label"]
    )
    for q in source.snapshot.questions:
        append(sheet, ["question", q.questionId, q.title, q.kind, None, None])
        choices = q.matrixRows[0].counts if q.matrixRows else q.counts
        for choice in choices:
            append(
                sheet,
                [
                    "choice",
                    q.questionId,
                    q.title,
                    q.kind,
                    json_value(choice.value),
                    choice.label,
                ],
            )
        for row in q.matrixRows:
            append(
                sheet,
                [
                    "matrix_row",
                    q.questionId,
                    q.title,
                    q.kind,
                    json_value(row.rowId),
                    row.label,
                ],
            )


def analysis_sheets(workbook: Any, source: SurveyExportSource) -> None:
    metrics = workbook.create_sheet("Metrics")
    append(
        metrics,
        [
            "question_id",
            "title",
            "kind",
            "relevant",
            "answered",
            "unanswered",
            "hidden",
            "invalid",
            "sum",
            "mean",
            "minimum",
            "maximum",
            "nps",
        ],
    )
    distributions = workbook.create_sheet("Distributions")
    append(
        distributions,
        [
            "question_id",
            "title",
            "row_id",
            "row_label",
            "choice_value_json",
            "choice_label",
            "count",
            "denominator",
            "percentage",
        ],
    )
    matrix = workbook.create_sheet("Matrix rows")
    append(
        matrix,
        ["question_id", "row_id", "row_label", "answered", "unanswered", "invalid"],
    )
    charts = workbook.create_sheet("Charts")
    append(
        charts,
        [
            "Charts use the same percentages and valid-answer denominators as Distributions."
        ],
    )
    anchor = 3
    for q in source.snapshot.questions:
        append(
            metrics,
            [
                q.questionId,
                q.title,
                q.kind,
                q.relevant,
                q.answered,
                q.unanswered,
                q.hidden,
                q.invalid,
                q.sum,
                q.mean,
                q.minimum,
                q.maximum,
                q.nps,
            ],
        )
        groups: list[tuple[str | None, str | None, int, list[AggregateCount]]] = (
            [(None, None, q.answered, q.counts)] if q.counts else []
        )
        for row in q.matrixRows:
            append(
                matrix,
                [
                    q.questionId,
                    row.rowId,
                    row.label,
                    row.answered,
                    row.unanswered,
                    row.invalid,
                ],
            )
            groups.append((row.rowId, row.label, row.answered, row.counts))
        for row_id, row_label, denominator, counts in groups:
            start = distributions.max_row + 1
            for bucket in counts:
                append(
                    distributions,
                    [
                        q.questionId,
                        q.title,
                        row_id,
                        row_label,
                        json_value(bucket.value),
                        bucket.label,
                        bucket.count,
                        denominator,
                        bucket.percentage,
                    ],
                )
            if not counts:
                continue
            chart = BarChart()
            chart.type = "bar"
            chart.title = excel_text(q.title + (f" / {row_label}" if row_label else ""))
            chart.y_axis.title = "Share (%)"
            chart.y_axis.scaling.min = 0
            chart.y_axis.scaling.max = 100
            chart.legend = None
            chart.width = 24
            chart.height = max(8, min(24, len(counts) * 0.65 + 3))
            chart.add_data(
                Reference(
                    distributions,
                    min_col=9,
                    min_row=start,
                    max_row=distributions.max_row,
                )
            )
            chart.set_categories(
                Reference(
                    distributions,
                    min_col=6,
                    min_row=start,
                    max_row=distributions.max_row,
                )
            )
            charts.add_chart(chart, f"A{anchor}")
            anchor += math.ceil(chart.height * 2.2) + 4
    pages = workbook.create_sheet("Last page")
    append(
        pages,
        ["page_id", "page_title", "count", "selected_participations", "percentage"],
    )
    for page in source.snapshot.lastPageCounts:
        append(
            pages,
            [
                page.pageId,
                page.title,
                page.count,
                source.snapshot.participationCount,
                page.count / source.snapshot.participationCount * 100
                if source.snapshot.participationCount
                else None,
            ],
        )


def workbook_bytes(workbook: Any, timestamp: str) -> bytes:
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.print_title_rows = "1:1"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="00338D")
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = (
                28 if column != 2 else 42
            )
    raw = io.BytesIO()
    workbook.save(raw)
    # Normalize both ZIP timestamps and openpyxl's save-time modified property.
    result = io.BytesIO()
    with (
        zipfile.ZipFile(raw) as source,
        zipfile.ZipFile(
            result, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as target,
    ):
        for name in sorted(source.namelist()):
            content = source.read(name)
            if name == "docProps/core.xml":
                for prefix, namespace in {
                    "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
                    "dc": "http://purl.org/dc/elements/1.1/",
                    "dcterms": "http://purl.org/dc/terms/",
                    "xsi": "http://www.w3.org/2001/XMLSchema-instance",
                }.items():
                    ElementTree.register_namespace(prefix, namespace)
                root = ElementTree.fromstring(content)
                for element in root:
                    if element.tag.endswith(("}created", "}modified")):
                        element.text = timestamp
                content = ElementTree.tostring(root, encoding="utf-8")
            entry = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(entry, content)
    return result.getvalue()


def render_tabular(
    product: ExportProduct, source: SurveyExportSource
) -> SurveyExportArtifact:
    s = AnalysisSnapshot.model_validate(source.snapshot.model_dump())
    info = metadata(source)
    if product == "responses_csv":
        columns, rows = response_rows(source)
        output = io.StringIO(newline="")
        writer = csv.writer(output, lineterminator="\r\n")
        writer.writerow(["record_type", *info, *columns])
        # A rectangular metadata record also preserves filter context for an empty export.
        writer.writerow(
            [
                "metadata",
                *[csv_value(v) for v in info.values()],
                *[None for _ in columns],
            ]
        )
        for row in rows:
            writer.writerow(
                [
                    "response",
                    *[csv_value(v) for v in info.values()],
                    *[csv_value(v) for v in row],
                ]
            )
        return SurveyExportArtifact(
            export_filename(product, s.id),
            "text/csv; charset=utf-8",
            RENDER_VERSION,
            output.getvalue().encode("utf-8-sig"),
        )
    if product not in {"responses_xlsx", "analysis_xlsx"}:
        raise SurveyExportRenderError("unsupported_tabular_product")
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "LeonAid Surveys"
    workbook.properties.title = excel_text(source.title)
    stamp = datetime.fromisoformat(s.createdAt)
    workbook.properties.created = stamp.replace(tzinfo=None)
    sheet = workbook.create_sheet("Metadata")
    append(sheet, ["field", "value"])
    for key, value in info.items():
        append(sheet, [key, value])
    catalogue(workbook, source)
    if product == "responses_xlsx":
        columns, rows = response_rows(source)
        sheet = workbook.create_sheet("Responses")
        append(sheet, columns)
        for row in rows:
            append(sheet, row)
    else:
        analysis_sheets(workbook, source)
    content = workbook_bytes(workbook, stamp.isoformat())
    return SurveyExportArtifact(
        export_filename(product, s.id), XLSX_MEDIA, RENDER_VERSION, content
    )
