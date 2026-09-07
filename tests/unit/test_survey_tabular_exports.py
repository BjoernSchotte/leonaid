"""Independent artifact parsing, golden values and adversarial export strings."""

import csv
import io
import json
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace
from uuid import NAMESPACE_URL, uuid5

import pytest
from openpyxl import load_workbook

from leonaid.adapters.survey_tabular_exports import render_tabular
from leonaid.application.surveys.analysis_snapshot import AnalysisSnapshot
from leonaid.application.surveys.export_rendering import (
    SurveyExportRenderError,
    SurveyExportSource,
)
from leonaid.application.surveys.response_selection import IndividualResponse

FIXTURES = Path(__file__).parents[1] / "fixtures" / "surveys"


def source() -> SurveyExportSource:
    snapshot = AnalysisSnapshot.model_validate_json(
        (FIXTURES / "export-snapshot.json").read_text()
    )
    answers = json.loads((FIXTURES / "analysis-golden.json").read_text())["responses"]
    return SurveyExportSource(
        "Rückmeldung – Golf & Krapfentaxi",
        snapshot,
        tuple(
            IndividualResponse(
                participationId=str(uuid5(NAMESPACE_URL, f"synthetic-export/{i}")),
                revision=1,
                status="completed" if i == 0 else "partial",
                answers=answer,
                currentPage="main",
                createdAt=snapshot.createdAt,
                completedAt=snapshot.createdAt if i == 0 else None,
            )
            for i, answer in enumerate(answers)
        ),
    )


def workbook(artifact):
    return load_workbook(io.BytesIO(artifact.content), data_only=False)


def rows_by_id(sheet):
    values = list(sheet.values)
    return {r[0]: dict(zip(values[0], r, strict=True)) for r in values[1:]}


def test_raw_csv_and_xlsx_roundtrip_structures_metadata_and_missing_values():
    data = source()
    csv_artifact = render_tabular("responses_csv", data)
    parsed = list(csv.DictReader(io.StringIO(csv_artifact.content.decode("utf-8-sig"))))
    assert parsed[0]["record_type"] == "metadata"
    assert parsed[0]["selected_participations"] == "5"
    assert json.loads(parsed[0]["statuses_json"]) == ["partial", "completed"]
    csv_rows = {
        row["participation_id"]: row
        for row in parsed
        if row["record_type"] == "response"
    }
    xlsx = workbook(render_tabular("responses_xlsx", data))
    xlsx_rows = rows_by_id(xlsx["Responses"])
    assert (
        set(csv_rows) == set(xlsx_rows) == {r.participationId for r in data.responses}
    )
    for response in data.responses:
        for key, answer in response.answers.items():
            value = xlsx_rows[response.participationId][f"q:{key}"]
            if isinstance(answer, (dict, list)):
                assert json.loads(value) == answer
                assert (
                    json.loads(csv_rows[response.participationId][f"q:{key}"]) == answer
                )
            else:
                assert ("" if answer == "" and value is None else value) == answer
                assert xlsx_rows[response.participationId][f"q:{key}:type"] == (
                    "string" if isinstance(answer, str) else "number"
                )
    first = xlsx_rows[data.responses[0].participationId]
    assert first["q:matrix:row:r1"] == "good"
    assert first["q:matrix:row:r2"] == "bad"
    missing = xlsx_rows[data.responses[-1].participationId]
    assert missing["q:nps"] is None
    info = dict(list(xlsx["Metadata"].values)[1:])
    assert info["snapshot_id"] == data.snapshot.id
    assert info["survey_title"] == data.title
    assert info["test_data"] is False
    assert xlsx["Questions"].max_row > len(data.snapshot.questions)


@pytest.mark.parametrize(
    "text",
    [
        "=1+1",
        "+SUM(1,2)",
        "-1+2",
        "@SUM(1,2)",
        ' \t=HYPERLINK("https://example.invalid")',
        "\ufeff=1",
        "\n=1",
        "Grüße 日本語\nZweite Zeile",
        "'=literal",
    ],
)
def test_formula_text_is_inert_and_unicode_is_retained(text):
    data = source()
    first = data.responses[0].model_copy(update={"answers": {"text": text}})
    data = SurveyExportSource(data.title, data.snapshot, (first, *data.responses[1:]))
    xlsx = workbook(render_tabular("responses_xlsx", data))
    sheet = xlsx["Responses"]
    headers = [c.value for c in sheet[1]]
    row = next(
        row
        for row in sheet.iter_rows(min_row=2)
        if row[0].value == first.participationId
    )
    cell = row[headers.index("q:text")]
    assert cell.value == text and cell.data_type == "s"
    assert all(
        c.data_type != "f" and c.hyperlink is None
        for s in xlsx
        for cells in s
        for c in cells
    )
    records = list(
        csv.DictReader(
            io.StringIO(
                render_tabular("responses_csv", data).content.decode("utf-8-sig")
            )
        )
    )
    record = next(r for r in records if r["participation_id"] == first.participationId)
    assert record["q:text"] == (text if text.startswith("Grüße") else "'" + text)


def test_analysis_workbook_contains_golden_metrics_and_charts_without_raw_text():
    data = source()
    artifact = render_tabular("analysis_xlsx", data)
    xlsx = workbook(artifact)
    metrics = rows_by_id(xlsx["Metrics"])
    assert metrics["nps"]["nps"] == pytest.approx(100 / 3)
    assert metrics["number"]["mean"] == 15
    assert metrics["follow"]["hidden"] == 2
    rows = list(xlsx["Distributions"].values)[1:]
    multi = [r for r in rows if r[0] == "multi"]
    assert [(r[6], r[7], r[8]) for r in multi] == [(2, 2, 100), (1, 2, 50)]
    matrix = [r for r in rows if r[0] == "matrix"]
    assert [(r[2], r[7], r[8]) for r in matrix] == [
        ("r1", 2, 50),
        ("r1", 2, 50),
        ("r2", 1, 0),
        ("r2", 1, 100),
    ]
    assert len(xlsx["Charts"]._charts) == 5
    with zipfile.ZipFile(io.BytesIO(artifact.content)) as archive:
        text = "\n".join(
            archive.read(name).decode()
            for name in archive.namelist()
            if name.endswith(".xml")
        )
        assert "SENSITIVE_" not in text
        assert "private_responses" not in text and "participationId" not in text
        assert "Distributions" in text and "dcterms:W3CDTF" in text


def test_rendered_bytes_are_deterministic_and_inputs_unchanged():
    data = source()
    before = data.snapshot.model_dump_json()
    for product in ("responses_csv", "responses_xlsx", "analysis_xlsx"):
        first = render_tabular(product, data)
        second = render_tabular(product, data)
        assert first.content == second.content and first.sha256 == second.sha256
    assert data.snapshot.model_dump_json() == before


@pytest.mark.parametrize("text", ["x" * 32768, "unrepresentable\x00control"])
def test_xlsx_never_silently_truncates_or_discards_unsupported_characters(text):
    data = source()
    first = data.responses[0].model_copy(update={"answers": {"text": text}})
    data = SurveyExportSource(data.title, data.snapshot, (first, *data.responses[1:]))
    with pytest.raises(SurveyExportRenderError, match="cell_text_unrepresentable"):
        render_tabular("responses_xlsx", data)


def test_empty_selection_keeps_metadata_and_raw_scope_is_required():
    data = source()
    payload = data.snapshot.model_dump()
    payload.update(
        participationCount=0,
        statusCounts={"in_progress": 0, "partial": 0, "completed": 0},
        questions=[],
        lastPageCounts=[],
    )
    empty = SurveyExportSource(data.title, AnalysisSnapshot.model_validate(payload), ())
    records = list(
        csv.DictReader(
            io.StringIO(
                render_tabular("responses_csv", empty).content.decode("utf-8-sig")
            )
        )
    )
    assert len(records) == 1 and records[0]["record_type"] == "metadata"
    assert records[0]["snapshot_id"] == data.snapshot.id
    assert records[0]["selected_participations"] == "0"
    assert workbook(render_tabular("responses_xlsx", empty))["Responses"].max_row == 1
    with pytest.raises(SurveyExportRenderError, match="raw_selection_required"):
        render_tabular("responses_csv", SurveyExportSource(data.title, data.snapshot))


@pytest.mark.parametrize("empty", [False, True])
def test_analysis_without_distributions_has_a_valid_printable_empty_state(empty):
    data = source()
    payload = data.snapshot.model_dump()
    payload["questions"] = (
        []
        if empty
        else [
            question
            for question in payload["questions"]
            if question["kind"] == "comment"
        ]
    )
    assert empty or payload["questions"]
    report = workbook(
        render_tabular(
            "analysis_xlsx",
            SurveyExportSource(data.title, AnalysisSnapshot.model_validate(payload)),
        )
    )
    charts = report["Charts"]
    assert not charts._charts
    assert charts["A2"].value == "No response distributions for this selection."
    assert "$A$1:$D$2" in str(charts.print_area)
    assert report["Metrics"].max_row == 1 + len(payload["questions"])


def test_raw_types_preserve_empty_null_boolean_and_high_precision_numbers():
    data = source()
    first = data.responses[0].model_copy(
        update={
            "answers": {
                "number": 9007199254740993,
                "nps": 0.12345678901234568,
                "text": "",
                "date": None,
                "gate": True,
                "matrix": {},
                "multi": [],
            }
        }
    )
    second = data.responses[1].model_copy(
        update={"answers": {"number": 0, "nps": -2.5}}
    )
    data = SurveyExportSource(
        data.title, data.snapshot, (first, second, *data.responses[2:])
    )
    rows = rows_by_id(workbook(render_tabular("responses_xlsx", data))["Responses"])
    row = rows[first.participationId]
    assert (row["q:number"], row["q:number:type"]) == ("9007199254740993", "number")
    assert (row["q:nps"], row["q:nps:type"]) == ("0.12345678901234568", "number")
    assert (row["q:text"], row["q:text:type"]) == (None, "string")
    assert (row["q:date"], row["q:date:type"]) == (None, "null")
    assert (row["q:gate"], row["q:gate:type"]) == (True, "boolean")
    assert row["q:conditional:type"] == "missing"
    assert rows[second.participationId]["q:number"] == 0
    assert rows[second.participationId]["q:nps"] == -2.5
    parsed = list(
        csv.DictReader(
            io.StringIO(
                render_tabular("responses_csv", data).content.decode("utf-8-sig")
            )
        )
    )
    assert (
        next(r for r in parsed if r["participation_id"] == first.participationId)[
            "q:gate"
        ]
        == "true"
    )
    assert (
        next(r for r in parsed if r["participation_id"] == second.participationId)[
            "q:nps"
        ]
        == "-2.5"
    )


def test_xlsx_bytes_do_not_depend_on_the_writer_clock(monkeypatch):
    class FirstClock(datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime(2030, 1, 1, tzinfo=timezone.utc)

    class LaterClock(datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime(2040, 12, 31, tzinfo=timezone.utc)

    monkeypatch.setattr(
        "openpyxl.writer.excel.datetime",
        SimpleNamespace(datetime=FirstClock, timezone=timezone),
    )
    first = render_tabular("analysis_xlsx", source())
    monkeypatch.setattr(
        "openpyxl.writer.excel.datetime",
        SimpleNamespace(datetime=LaterClock, timezone=timezone),
    )
    second = render_tabular("analysis_xlsx", source())
    assert first.content == second.content


def test_malformed_raw_selection_cannot_be_exported():
    data = source()
    with pytest.raises(SurveyExportRenderError, match="duplicate_participation"):
        render_tabular(
            "responses_xlsx",
            SurveyExportSource(
                data.title, data.snapshot, (*data.responses[:-1], data.responses[0])
            ),
        )
    changed = data.responses[0].model_copy(update={"status": "in_progress"})
    with pytest.raises(
        SurveyExportRenderError, match="response_status_outside_selection"
    ):
        render_tabular(
            "responses_csv",
            SurveyExportSource(
                data.title, data.snapshot, (changed, *data.responses[1:])
            ),
        )
