"""Real API/outbox/worker/private-object survey export proof; synthetic data only."""

import asyncio
import csv
import hashlib
import io
import json
import os
import secrets
import sys
from datetime import datetime, timezone
from pathlib import Path
from uuid import UUID, uuid4

import asyncpg
import httpx
from openpyxl import load_workbook
from pypdf import PdfReader

from leonaid.domain.sessions import SESSION_LIFETIME, session_token_digest


async def main():
    conn = await asyncpg.connect(os.environ["CORE_DATABASE_URL"])
    token = Path("/proof/session.env").read_text().strip().split("=", 1)[1]
    admin = {"Cookie": f"__Host-leonaid_session={token}"}
    state_path = Path("/proof/export-state.json")
    async with httpx.AsyncClient(base_url="http://api:8000", timeout=60) as client:

        async def call(method, path, body=None, auth=admin, expected=200):
            response = await client.request(method, path, json=body, headers=auth)
            assert response.status_code == expected, (
                path,
                response.status_code,
                expected,
            )
            if expected == 200 and "/exports" in path:
                assert response.headers["cache-control"] == "no-store"
            return response

        if sys.argv[1] == "prepare":
            now = datetime.now(timezone.utc)
            fixture = json.loads(
                Path("tests/fixtures/surveys/analysis-golden.json").read_text()
            )
            sid = uuid4()
            path = f"/api/v1/surveys/{sid}"
            await call(
                "POST",
                path,
                {
                    "operationId": "create",
                    "title": "Synthetic export – Unicode ü",
                    "definition": fixture["definition"],
                },
            )
            version = (
                await call(
                    "POST",
                    path + "/publish",
                    {"operationId": "publish", "expectedRevision": 1},
                )
            ).json()
            for sample in fixture["responses"]:
                await conn.execute(
                    """INSERT INTO survey_participation(id,survey_id,version_id,resume_digest,status,answers,current_page,inactivity_timeout_seconds,created_at,last_answer_changed_at,completed_at)
                    VALUES($1,$2,$3,$4,'completed',$5::jsonb,'main',3600,$6,$6,$6)""",
                    uuid4(),
                    sid,
                    UUID(version["id"]),
                    hashlib.sha256(secrets.token_bytes(48)).hexdigest(),
                    json.dumps(sample),
                    now,
                )
            snapshot = (
                await call(
                    "POST",
                    path + "/analysis",
                    {"operationId": "snapshot", "filter": {"versionId": version["id"]}},
                )
            ).json()
            uid, member_token = uuid4(), secrets.token_urlsafe(48)
            await conn.execute(
                "INSERT INTO user_account(id,email,display_name,status,email_verified_at) VALUES($1,$2,'Synthetic export member','active',$3)",
                uid,
                f"{uid}@example.invalid",
                now,
            )
            await conn.execute(
                "INSERT INTO user_session(id,user_id,token_digest,expires_at,last_seen_at,fresh_login_at,created_at,updated_at) VALUES($1,$2,$3,$4,$5,$5,$5,$5)",
                uuid4(),
                uid,
                session_token_digest(member_token),
                now + SESSION_LIFETIME,
                now,
            )
            await conn.execute(
                "INSERT INTO survey_grant(survey_id,user_id,capability) VALUES($1,$2,'export_reports')",
                sid,
                uid,
            )
            member = {"Cookie": f"__Host-leonaid_session={member_token}"}
            jobs = {}
            for product in (
                "responses_csv",
                "responses_xlsx",
                "analysis_xlsx",
                "analysis_pdf",
            ):
                body = {
                    "operationId": product,
                    "snapshotId": snapshot["id"],
                    "product": product,
                }
                job = (await call("POST", path + "/exports", body)).json()
                assert job["status"] == "queued"
                assert (await call("POST", path + "/exports", body)).json()[
                    "id"
                ] == job["id"]
                await call(
                    "POST",
                    path + "/exports",
                    {
                        **body,
                        "product": "responses_csv"
                        if product == "analysis_pdf"
                        else "analysis_pdf",
                    },
                    expected=409,
                )
                await call("GET", path + f"/exports/{job['id']}/download", expected=409)
                jobs[product] = job["id"]
                assert not any(
                    key in job for key in ("bucket", "object_key", "object_version")
                )
            await call(
                "POST",
                path + "/exports",
                {
                    "operationId": "denied-raw",
                    "snapshotId": snapshot["id"],
                    "product": "responses_csv",
                },
                member,
                404,
            )
            member_job = (
                await call(
                    "POST",
                    path + "/exports",
                    {
                        "operationId": "revoked-before-processing",
                        "snapshotId": snapshot["id"],
                        "product": "analysis_xlsx",
                    },
                    member,
                )
            ).json()
            await conn.execute(
                "DELETE FROM survey_grant WHERE survey_id=$1 AND user_id=$2", sid, uid
            )
            await call(
                "GET", path + f"/exports/{member_job['id']}", auth=member, expected=404
            )
            await call(
                "POST",
                path + "/exports",
                {
                    "operationId": "foreign",
                    "snapshotId": str(uuid4()),
                    "product": "analysis_xlsx",
                },
                expected=404,
            )
            await call(
                "GET", path + f"/exports/{jobs['responses_csv']}", auth={}, expected=401
            )
            assert (
                await conn.fetchval(
                    "SELECT count(*) FROM survey_export_job WHERE survey_id=$1", sid
                )
                == 5
            )
            try:
                await conn.execute(
                    "UPDATE survey_export_job SET title='changed' WHERE id=$1",
                    UUID(jobs["responses_csv"]),
                )
                raise AssertionError("Export input mutation succeeded")
            except asyncpg.IntegrityConstraintViolationError:
                pass
            missing_font_path = f"/api/v1/surveys/{uuid4()}"
            await call(
                "POST",
                missing_font_path,
                {
                    "operationId": "create",
                    "title": "日本語",
                    "definition": fixture["definition"],
                },
            )
            missing_font_version = (
                await call(
                    "POST",
                    missing_font_path + "/publish",
                    {"operationId": "publish", "expectedRevision": 1},
                )
            ).json()
            missing_font_snapshot = (
                await call(
                    "POST",
                    missing_font_path + "/analysis",
                    {
                        "operationId": "snapshot",
                        "filter": {"versionId": missing_font_version["id"]},
                    },
                )
            ).json()
            missing_font_job = (
                await call(
                    "POST",
                    missing_font_path + "/exports",
                    {
                        "operationId": "missing-font",
                        "snapshotId": missing_font_snapshot["id"],
                        "product": "analysis_pdf",
                    },
                )
            ).json()
            state_path.write_text(
                json.dumps(
                    {
                        "path": path,
                        "surveyId": str(sid),
                        "snapshot": snapshot,
                        "jobs": jobs,
                        "memberId": str(uid),
                        "memberToken": member_token,
                        "cancelledJob": member_job["id"],
                        "missingFontPath": missing_font_path,
                        "missingFontJob": missing_font_job["id"],
                    }
                )
            )
            print(
                "PASS: queued all four exports, atomic idempotency, immutable inputs, permission and unauthenticated denial; worker is stopped"
            )
        else:
            state = json.loads(state_path.read_text())
            path = state["path"]
            for product, jid in state["jobs"].items():
                for _ in range(200):
                    job = (await call("GET", path + f"/exports/{jid}")).json()
                    if job["status"] in {"available", "failed"}:
                        break
                    await asyncio.sleep(0.2)
                assert job["status"] == "available", (
                    product,
                    job["status"],
                    job["errorCode"],
                )
                response = await call("GET", path + f"/exports/{jid}/download")
                assert response.headers["content-disposition"].startswith("attachment;")
                assert response.headers["x-content-type-options"] == "nosniff"
                content = response.content
                assert len(content) == job["sizeBytes"]
                assert (
                    token.encode() not in content
                    and state["memberToken"].encode() not in content
                )
                if product == "responses_csv":
                    records = list(
                        csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
                    )
                    assert records[0]["record_type"] == "metadata"
                    assert (
                        len([r for r in records if r["record_type"] == "response"])
                        == state["snapshot"]["participationCount"]
                    )
                    assert records[0]["snapshot_id"] == state["snapshot"]["id"]
                elif product == "analysis_pdf":
                    reader = PdfReader(io.BytesIO(content))
                    text = "\n".join(page.extract_text() for page in reader.pages)
                    assert (
                        state["snapshot"]["id"] in text
                        and state["snapshot"]["filter"]["versionId"] in text
                    )
                    assert "33,33" in text and "Mittelwert" in text
                    assert "SENSITIVE_" not in text and len(reader.pages) >= 2
                    Path("/proof/survey-worker-report.pdf").write_bytes(content)
                else:
                    workbook = load_workbook(io.BytesIO(content), data_only=False)
                    meta = dict(list(workbook["Metadata"].values)[1:])
                    assert meta["snapshot_id"] == state["snapshot"]["id"]
                    if product == "responses_xlsx":
                        assert (
                            workbook["Responses"].max_row
                            == state["snapshot"]["participationCount"] + 1
                        )
                    else:
                        rows = list(workbook["Metrics"].values)
                        metrics = {r[0]: dict(zip(rows[0], r)) for r in rows[1:]}
                        for question in state["snapshot"]["questions"]:
                            for key in (
                                "relevant",
                                "answered",
                                "unanswered",
                                "hidden",
                                "invalid",
                                "mean",
                                "nps",
                            ):
                                expected = question[key]
                                actual = metrics[question["questionId"]][key]
                                assert (
                                    actual == expected
                                    or isinstance(actual, float)
                                    and abs(actual - expected) < 1e-10
                                )
                        assert len(workbook["Charts"]._charts) > 0
                stored = await conn.fetchrow(
                    "SELECT bucket,object_key,object_version,sha256 FROM survey_export_job WHERE id=$1",
                    UUID(jid),
                )
                assert (
                    stored["object_version"]
                    and stored["sha256"] == hashlib.sha256(content).hexdigest()
                )
                anonymous = await client.get(
                    os.environ["OBJECT_STORAGE_ENDPOINT_URL"].rstrip("/")
                    + f"/{stored['bucket']}/{stored['object_key']}"
                )
                assert anonymous.status_code in (401, 403)
            cancelled = await conn.fetchrow(
                "SELECT status,bucket FROM survey_export_job WHERE id=$1",
                UUID(state["cancelledJob"]),
            )
            assert dict(cancelled) == {"status": "cancelled", "bucket": None}
            for _ in range(100):
                failed = (
                    await call(
                        "GET",
                        state["missingFontPath"]
                        + f"/exports/{state['missingFontJob']}",
                    )
                ).json()
                if failed["status"] in {"retrying", "failed"}:
                    break
                await asyncio.sleep(0.2)
            assert (
                failed["status"] in {"retrying", "failed"}
                and failed["errorCode"] == "survey_export_failed"
            )
            await call(
                "GET",
                state["missingFontPath"]
                + f"/exports/{state['missingFontJob']}/download",
                expected=409,
            )
            failed_record = await conn.fetchrow(
                "SELECT j.bucket, e.last_error_detail FROM survey_export_job j JOIN outbox_event e ON e.id=j.event_id WHERE j.id=$1",
                UUID(state["missingFontJob"]),
            )
            assert dict(failed_record) == {
                "bucket": None,
                "last_error_detail": "survey_export_failed",
            }
            assert await conn.fetchval(
                "SELECT bool_and(payload='{}'::jsonb) FROM outbox_event WHERE event_type='survey.export.render.v1'"
            )
            # Existing artifacts become inaccessible immediately when the survey is trashed.
            summary = (await call("GET", path)).json()
            await call(
                "POST",
                path + "/transition",
                {
                    "operationId": "trash-after-export",
                    "expectedRevision": summary["revision"],
                    "action": "trash",
                },
            )
            for jid in state["jobs"].values():
                await call("GET", path + f"/exports/{jid}", expected=404)
                await call("GET", path + f"/exports/{jid}/download", expected=404)
            # Retention cannot lose object references by deleting the survey first.
            try:
                await conn.execute(
                    "DELETE FROM survey WHERE id=$1", UUID(state["surveyId"])
                )
                raise AssertionError("Referenced survey deletion succeeded")
            except asyncpg.ForeignKeyViolationError:
                pass
            Path("/proof/survey-exports-proof.json").write_text(
                json.dumps(
                    {
                        "products": list(state["jobs"]),
                        "realWorker": True,
                        "parsedDownloads": 4,
                        "anonymousObjectAccessDenied": True,
                        "revokedQueuedJobCancelled": True,
                        "deletedSurveyDownloadsDenied": True,
                        "outboxPayloadsContainNoAnswers": True,
                        "objectReferencesRetained": True,
                        "missingGlyphJobFailedWithoutArtifact": True,
                    }
                )
            )
            state_path.unlink()
            print(
                "PASS: all four real worker/private-storage downloads parsed against snapshot; cancellation, deletion, anonymous storage and orphan prevention proven"
            )
    await conn.close()


asyncio.run(main())
