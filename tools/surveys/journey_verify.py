"""Inspect the real browser downloads and the final SQL/object erasure state."""

import asyncio
import csv
import io
import json
import os
from pathlib import Path
from uuid import UUID

import asyncpg
from openpyxl import load_workbook
from pypdf import PdfReader

from leonaid.adapters.storage.s3 import S3ObjectStorage

PROOF = Path("/proof")


def exports(state, stem):
    snapshot = state["snapshot"]
    assert snapshot["versionNumber"] == 1 and snapshot["participationCount"] == 1
    assert snapshot["statusCounts"] == {"completed": 1, "partial": 0, "in_progress": 0}
    metrics = {q["questionId"]: q for q in snapshot["questions"]}
    if state["template"] == "krapfentaxi":
        assert metrics["delivery_rating"]["mean"] == 5
        assert metrics["nps"]["nps"] == 100
    else:
        assert state["answers"]["event_rating"] == {
            "organization": 3,
            "course": 3,
            "catering": 3,
        }
    csv_rows = list(
        csv.DictReader(
            io.StringIO(
                (PROOF / f"{stem}-responses_csv.csv").read_text(encoding="utf-8-sig")
            )
        )
    )
    assert csv_rows[0]["snapshot_id"] == snapshot["id"]
    raw = [r for r in csv_rows if r["record_type"] == "response"]
    assert len(raw) == 1 and raw[0]["participation_id"] == state["pid"]
    assert raw[0]["status"] == "completed"
    comment = "food_feedback" if state["template"] == "golf" else "notes"
    assert raw[0][f"q:{comment}"] == state["answers"][comment]
    rows = list(
        load_workbook(PROOF / f"{stem}-responses_xlsx.xlsx")["Responses"].values
    )
    assert len(rows) == 2
    xlsx_raw = dict(zip(rows[0], rows[1]))
    for key, value in xlsx_raw.items():
        assert str(value if value is not None else "") == raw[0][key], key
    book = load_workbook(PROOF / f"{stem}-analysis_xlsx.xlsx")
    assert dict(list(book["Metadata"].values)[1:])["snapshot_id"] == snapshot["id"]
    rows = list(book["Metrics"].values)
    assert len(rows) == len(metrics) + 1
    for row in rows[1:]:
        actual = dict(zip(rows[0], row))
        expected = metrics[row[0]]
        for key in (
            "relevant",
            "answered",
            "unanswered",
            "hidden",
            "invalid",
            "mean",
            "minimum",
            "maximum",
            "nps",
        ):
            assert actual[key] == expected[key], (row[0], key)
    assert book["Charts"]._charts
    text = "\n".join(
        p.extract_text() for p in PdfReader(PROOF / f"{stem}-analysis_pdf.pdf").pages
    )
    assert snapshot["id"] in text
    for index, question in enumerate(snapshot["questions"], 1):
        counts = " ".join(
            str(question[key])
            for key in ("relevant", "answered", "unanswered", "hidden", "invalid")
        )
        assert (
            f"{index}. {question['title']}\nRelevant Beantwortet Unbeantwortet Ausgeblendet Ungültig\n{counts}"
            in text
        )
    assert state["answers"][comment] not in text
    assert state["answers"][comment] not in " ".join(
        str(c) for sheet in book for row in sheet.values for c in row
    )


async def main():
    conn = await asyncpg.connect(os.environ["CORE_DATABASE_URL"])
    storage = S3ObjectStorage(
        endpoint_url=os.environ["OBJECT_STORAGE_ENDPOINT_URL"],
        access_key=os.environ["OBJECT_STORAGE_ACCESS_KEY"],
        secret_key=os.environ["OBJECT_STORAGE_SECRET_KEY"],
        bucket=os.environ["OBJECT_STORAGE_BUCKET"],
        region=os.environ.get("OBJECT_STORAGE_REGION", "us-east-1"),
    )
    results = []
    try:
        for template in ["krapfentaxi", "golf"]:
            for width in [1440, 390]:
                stem = f"journey-{template}-{width}"
                state = json.loads((PROOF / f"{stem}.json").read_text())
                assert state["template"] == template and state["width"] == width
                assert all(
                    state[k]
                    for k in [
                        "partialResumeAndVersionIsolation",
                        "outsiderDenied",
                        "erasureCompleted",
                    ]
                )
                exports(state, stem)
                sid = UUID(state["sid"])
                for table, column in [
                    ("survey", "id"),
                    ("survey_draft", "survey_id"),
                    ("survey_version", "survey_id"),
                    ("survey_participation", "survey_id"),
                    ("survey_operation", "survey_id"),
                    ("survey_invitation", "survey_id"),
                    ("survey_analysis_snapshot", "survey_id"),
                    ("survey_export_job", "survey_id"),
                    ("survey_grant", "survey_id"),
                ]:
                    assert not await conn.fetchval(
                        f"SELECT EXISTS(SELECT 1 FROM {table} WHERE {column}=$1)", sid
                    ), table
                assert await conn.fetchval(
                    "SELECT completed_at IS NOT NULL FROM survey_deletion WHERE survey_id=$1",
                    sid,
                )
                # Read every S3 version under this owned survey prefix, not merely
                # the latest version: a delete marker must not conceal retained data.
                paginator = storage._client.get_paginator("list_object_versions")
                for objects in paginator.paginate(
                    Bucket=storage.bucket, Prefix=f"surveys/{sid}/"
                ):
                    assert not objects.get("Versions") and not objects.get(
                        "DeleteMarkers"
                    )
                results.append(
                    {
                        "template": template,
                        "viewport": [width, 960],
                        "partialResumeAndVersionIsolation": True,
                        "outsiderDenied": True,
                        "browserExportsParsed": 4,
                        "csvXlsxAgreement": True,
                        "snapshotMetricsAgree": True,
                        "sqlContentErased": True,
                        "allObjectVersionsErased": True,
                    }
                )
    finally:
        await conn.close()
    (PROOF / "journeys-proof.json").write_text(
        json.dumps({"syntheticOnly": True, "journeys": results}, indent=2) + "\n"
    )
    print(
        "PASS: four complete browser journeys; 16 real downloads parsed; SQL content and every object version erased"
    )


asyncio.run(main())
