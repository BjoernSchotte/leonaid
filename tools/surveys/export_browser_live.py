"""Synthetic browser fixtures and independent parsing of actual downloaded exports."""

import asyncio
import csv
import hashlib
import io
import json
import os
import secrets
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from uuid import UUID, uuid4

import asyncpg
import httpx
from openpyxl import load_workbook
from pypdf import PdfReader

from leonaid.domain.sessions import SESSION_LIFETIME, session_token_digest

PROOF = Path("/proof")


async def main():
    if sys.argv[1] == "verify-revoke-when-ready":
        deadline = time.monotonic() + 180
        while not (PROOF / "export-revoke-request").exists():
            if time.monotonic() >= deadline:
                raise TimeoutError("Browser did not request export verification")
            await asyncio.sleep(0.1)
    conn = await asyncpg.connect(os.environ["CORE_DATABASE_URL"])
    state_path = PROOF / "export-browser-access.json"
    if sys.argv[1] == "seed":
        now = datetime.now(timezone.utc)
        fixture = json.loads(
            Path("tests/fixtures/surveys/analysis-golden.json").read_text()
        )
        admin_token = (PROOF / "session.env").read_text().strip().split("=", 1)[1]
        sid, uid, member_token = uuid4(), uuid4(), secrets.token_urlsafe(48)
        async with httpx.AsyncClient(
            base_url="http://api:8000",
            headers={"Cookie": f"__Host-leonaid_session={admin_token}"},
        ) as client:
            created = await client.post(
                f"/api/v1/surveys/{sid}",
                json={
                    "operationId": "create",
                    "title": "Synthetic populated browser export",
                    "definition": fixture["definition"],
                },
            )
            assert created.status_code == 200
            published = await client.post(
                f"/api/v1/surveys/{sid}/publish",
                json={"operationId": "publish", "expectedRevision": 1},
            )
            assert published.status_code == 200
            version = UUID(published.json()["id"])
        for index, answers in enumerate(fixture["responses"]):
            await conn.execute(
                """INSERT INTO survey_participation(id,survey_id,version_id,resume_digest,status,answers,current_page,inactivity_timeout_seconds,created_at,last_answer_changed_at,completed_at)
                VALUES($1,$2,$3,$4,$5,$6::jsonb,'main',3600,$7,$7,$8)""",
                uuid4(),
                sid,
                version,
                hashlib.sha256(secrets.token_bytes(48)).hexdigest(),
                "partial" if index == 4 else "completed",
                json.dumps(answers),
                now,
                None if index == 4 else now,
            )
        await conn.execute(
            "INSERT INTO user_account(id,email,display_name,status,email_verified_at) VALUES($1,$2,'Synthetic report exporter','active',$3)",
            uid,
            f"{uid}@example.invalid",
            now,
        )
        await conn.execute(
            "INSERT INTO user_session(id,user_id,token_digest,expires_at,last_seen_at,fresh_login_at,created_at,updated_at) VALUES($1,$2,$3,$4,$5,$5,$5,$5)",
            uuid4(),
            uid,
            session_token_digest(member_token),
            now + SESSION_LIFETIME,
            now,
        )
        for capability in ("view_aggregates", "export_reports"):
            await conn.execute(
                "INSERT INTO survey_grant(survey_id,user_id,capability) VALUES($1,$2,$3)",
                sid,
                uid,
                capability,
            )
        export_only = {}
        for capability in ("export_raw", "export_reports"):
            export_uid, token = uuid4(), secrets.token_urlsafe(48)
            await conn.execute(
                "INSERT INTO user_account(id,email,display_name,status,email_verified_at) VALUES($1,$2,'Synthetic export-only member','active',$3)",
                export_uid,
                f"{export_uid}@example.invalid",
                now,
            )
            await conn.execute(
                "INSERT INTO user_session(id,user_id,token_digest,expires_at,last_seen_at,fresh_login_at,created_at,updated_at) VALUES($1,$2,$3,$4,$5,$5,$5,$5)",
                uuid4(),
                export_uid,
                session_token_digest(token),
                now + SESSION_LIFETIME,
                now,
            )
            await conn.execute(
                "INSERT INTO survey_grant(survey_id,user_id,capability) VALUES($1,$2,$3)",
                sid,
                export_uid,
                capability,
            )
            export_only[capability] = token
        state_path.write_text(
            json.dumps(
                {
                    "surveyId": str(sid),
                    "memberId": str(uid),
                    "memberToken": member_token,
                    "exportOnly": export_only,
                }
            )
        )
        print(
            "PASS: five synthetic responses (four completed, one partial) and report-only persona seeded"
        )
    elif sys.argv[1] in {"verify-revoke", "verify-revoke-when-ready"}:
        state = json.loads(state_path.read_text())
        observed = json.loads((PROOF / "export-populated-browser.json").read_text())
        snapshot = observed["snapshot"]
        assert snapshot["participationCount"] == 5
        assert snapshot["statusCounts"] == {
            "completed": 4,
            "partial": 1,
            "in_progress": 0,
        }
        metrics = {q["questionId"]: q for q in snapshot["questions"]}
        assert abs(metrics["nps"]["nps"] - 100 / 3) < 1e-10
        assert metrics["number"]["mean"] == 15
        # These are files downloaded by the browser, never a second API download.
        csv_rows = list(
            csv.DictReader(
                io.StringIO(
                    (PROOF / "browser-responses_csv.csv").read_text(
                        encoding="utf-8-sig"
                    )
                )
            )
        )
        assert csv_rows[0]["snapshot_id"] == snapshot["id"]
        raw = [r for r in csv_rows if r["record_type"] == "response"]
        assert len(raw) == 5 and sum(r["status"] == "partial" for r in raw) == 1
        assert any(r["q:text"] == "SENSITIVE_FIRST_TEXT" for r in raw)
        raw_book = load_workbook(PROOF / "browser-responses_xlsx.xlsx")
        rows = list(raw_book["Responses"].values)
        raw_xlsx = [dict(zip(rows[0], row)) for row in rows[1:]]
        assert len(raw_xlsx) == 5
        assert {r["participation_id"] for r in raw_xlsx} == {
            r["participation_id"] for r in raw
        }
        for r in raw_xlsx:
            corresponding = next(
                c for c in raw if c["participation_id"] == r["participation_id"]
            )
            for key in ("q:text", "q:matrix", "q:multi", "q:text:type", "status"):
                assert (
                    str(r[key] if r[key] is not None else "") == corresponding[key]
                ), key
        book = load_workbook(PROOF / "browser-analysis_xlsx.xlsx")
        meta = dict(list(book["Metadata"].values)[1:])
        assert meta["snapshot_id"] == snapshot["id"]
        rows = list(book["Metrics"].values)
        for row in rows[1:]:
            actual = dict(zip(rows[0], row))
            expected = metrics[row[0]]
            for key in (
                "relevant",
                "answered",
                "unanswered",
                "hidden",
                "invalid",
                "mean",
                "minimum",
                "maximum",
                "nps",
            ):
                assert (
                    actual[key] == expected[key]
                    or isinstance(actual[key], float)
                    and abs(actual[key] - expected[key]) < 1e-10
                ), key
        text = "\n".join(
            p.extract_text()
            for p in PdfReader(PROOF / "browser-analysis_pdf.pdf").pages
        )
        assert snapshot["id"] in text and "33,33" in text and "Mittelwert" in text
        assert "Net Promoter Score 33,33" in text
        assert "Summe 30\nMittelwert 15\nMinimum 10\nMaximum 20" in text
        for index, question in enumerate(snapshot["questions"], 1):
            counts = " ".join(
                str(question[key])
                for key in ("relevant", "answered", "unanswered", "hidden", "invalid")
            )
            assert (
                f"{index}. {question['title']}\nRelevant Beantwortet Unbeantwortet Ausgeblendet Ungültig\n{counts}"
                in text
            )
        for expected in (
            "Alpha 2 2 100 %",
            "Beta 1 2 50 %",
            "good 1 2 50 %",
            "bad 1 2 50 %",
            "good 0 1 0 %",
            "bad 1 1 100 %",
        ):
            assert expected in text
        assert "SENSITIVE_" not in text
        assert "SENSITIVE_" not in " ".join(
            str(c) for sheet in book for row in sheet.values for c in row
        )
        assert len(book["Charts"]._charts) > 0
        await conn.execute(
            "DELETE FROM survey_grant WHERE survey_id=$1 AND user_id=$2 AND capability='export_reports'",
            UUID(state["surveyId"]),
            UUID(state["memberId"]),
        )
        assert (
            await conn.fetchval(
                "SELECT count(*) FROM survey_grant WHERE survey_id=$1 AND user_id=$2 AND capability='view_aggregates'",
                UUID(state["surveyId"]),
                UUID(state["memberId"]),
            )
            == 1
        )
        # Only a content-free summary becomes a publishable proof artifact.
        (PROOF / "export-browser-values-proof.json").write_text(
            json.dumps(
                {
                    "participations": 5,
                    "completed": 4,
                    "partial": 1,
                    "nps": 100 / 3,
                    "mean": 15,
                    "browserDownloadsParsed": 4,
                    "rawCsvXlsxAgreement": True,
                    "aggregateFilesExcludeRawText": True,
                    "reportGrantRevoked": True,
                },
                indent=2,
            )
        )
        print(
            "PASS: actual browser downloads match displayed golden snapshot; report permission revoked while aggregate access retained"
        )
    else:
        raise ValueError("Unknown phase")
    await conn.close()


if sys.argv[1] == "verify-revoke-when-ready":
    try:
        asyncio.run(main())
    except BaseException:
        result = "failed"
        raise
    else:
        result = "passed"
    finally:
        # Atomic, content-free handshake; private SQL credentials stay in this container.
        temporary = PROOF / "export-revoke-result.tmp"
        temporary.write_text(result)
        temporary.replace(PROOF / "export-revoke-result")
else:
    asyncio.run(main())
